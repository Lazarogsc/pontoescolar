"""Revisão de texto dos materiais formatados (Word ou PDF), sem IA: LanguageTool.

O LanguageTool (código aberto, roda no nosso servidor — serviço "revisor" no compose)
aponta ortografia, gramática, concordância, crase e pontuação em português. Ele erra às
vezes (falso alarme e erro que passa), por isso nada é corrigido sozinho: o revisor aceita
ou recusa cada sugestão, e o sistema gera um Word com "Controlar alterações" só com o que
foi aceito, mantendo a formatação do diagramador. O revisor ainda abre no Word e confere.

- Word (.docx): corrige no próprio arquivo, parágrafo por parágrafo (inclui tabelas).
- PDF: não dá para editar no lugar; gera um Word com o texto extraído e as correções marcadas.
"""

import copy
import json
import os
import re
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader

from app import arquivos, estado, materiais
from app.util import agora

URL = os.getenv("LANGUAGETOOL_URL", "http://localhost:8010")
LOTE = 20_000  # caracteres por consulta ao LanguageTool
# Regras que só fazem barulho em material didático: enunciados "a) ...", listas sem ponto final.
IGNORAR_REGRAS = {"UPPERCASE_SENTENCE_START", "PUNCTUATION_PARAGRAPH_END"}
CHAVE_DICIONARIO = "dicionario_revisao"
PAGINAS_MAX = 200


class CorretorIndisponivel(RuntimeError):
    pass


# ---------- Texto do documento ----------


# O que cada elemento de um trecho (w:r) vale em texto. Análise e correção usam a mesma conta,
# então a posição de cada achado é a mesma na hora de marcar a alteração.
_VALOR = {qn("w:tab"): "\t", qn("w:br"): "\n", qn("w:cr"): "\n", qn("w:noBreakHyphen"): "-"}


def _texto_run(r) -> str:
    return "".join((x.text or "") if x.tag == qn("w:t") else _VALOR.get(x.tag, "") for x in r)


def _texto(p: Paragraph) -> str:
    """Texto dos trechos diretos do parágrafo (links e campos ficam de fora da revisão)."""
    return "".join(_texto_run(r) for r in p._p.iterchildren(qn("w:r")))


def _paragrafos_docx(doc) -> list[tuple[str, Paragraph]]:
    """Parágrafos na ordem do documento, com tabelas (célula mesclada aparece uma vez só)."""
    lista, vistos, tabela, paragrafo = [], set(), 0, 0
    corpo = doc.element.body
    for filho in corpo.iterchildren():
        if filho.tag == qn("w:p"):
            paragrafo += 1
            lista.append((f"Parágrafo {paragrafo}", Paragraph(filho, doc)))
        elif filho.tag == qn("w:tbl"):
            tabela += 1
            for li, linha in enumerate(Table(filho, doc).rows, start=1):
                for celula in linha.cells:
                    if id(celula._tc) in vistos:
                        continue
                    vistos.add(id(celula._tc))
                    for p in celula.paragraphs:
                        lista.append((f"Tabela {tabela}, linha {li}", p))
    return lista


def _juntar_linhas_pdf(texto: str) -> list[str]:
    """Junta as linhas quebradas do PDF em parágrafos (e desfaz a hifenização de fim de linha)."""
    paragrafos, atual = [], ""
    for linha in (l.strip() for l in texto.splitlines()):
        if not linha:
            if atual:
                paragrafos.append(atual)
            atual = ""
            continue
        if atual.endswith("-") and linha[:1].islower():
            atual = atual[:-1] + linha
        elif atual and not re.search(r"[.:;?!)]$", atual) and not re.match(r"^(\d+[.)]|[a-eA-E][.)]|[•\-–])\s", linha):
            atual += " " + linha
        else:
            if atual:
                paragrafos.append(atual)
            atual = linha
    if atual:
        paragrafos.append(atual)
    return paragrafos


def trechos(conteudo: bytes, ext: str) -> list[dict]:
    """[{rotulo, texto}] de um .docx ou .pdf."""
    if ext == ".docx":
        doc = Document(BytesIO(conteudo))
        return [{"rotulo": rotulo, "texto": _texto(p)} for rotulo, p in _paragrafos_docx(doc)]
    if ext == ".pdf":
        leitor = PdfReader(BytesIO(conteudo))
        if len(leitor.pages) > PAGINAS_MAX:
            raise ValueError(f"O PDF tem {len(leitor.pages)} páginas; a revisão aceita até {PAGINAS_MAX}.")
        lista = []
        for n, pagina in enumerate(leitor.pages, start=1):
            for paragrafo in _juntar_linhas_pdf(pagina.extract_text() or ""):
                lista.append({"rotulo": f"Página {n}", "texto": paragrafo})
        return lista
    raise ValueError("A revisão lê Word (.docx) ou PDF.")


# ---------- LanguageTool ----------


def _consultar(texto: str) -> list[dict]:
    corpo = urllib.parse.urlencode({"language": "pt-BR", "text": texto}).encode()
    try:
        with urllib.request.urlopen(urllib.request.Request(f"{URL}/v2/check", data=corpo), timeout=120) as r:
            return json.loads(r.read())["matches"]
    except (urllib.error.URLError, TimeoutError, KeyError, json.JSONDecodeError) as e:
        raise CorretorIndisponivel(
            "O corretor (LanguageTool) não respondeu. Confira se o serviço 'revisor' está no ar e tente de novo."
        ) from e


def dicionario() -> list[str]:
    return sorted(estado.ler("config").get(CHAVE_DICIONARIO, []))


def adicionar_ao_dicionario(palavra: str) -> list[str]:
    palavra = palavra.strip()
    if not palavra or len(palavra) > 60 or " " in palavra:
        raise ValueError("Informe uma palavra (sem espaços).")
    palavras = set(estado.ler("config").get(CHAVE_DICIONARIO, [])) | {palavra}
    estado.definir("config", CHAVE_DICIONARIO, sorted(palavras))
    return sorted(palavras)


def _contexto(texto: str, inicio: int, fim: int, folga: int = 45) -> dict:
    a, b = max(0, inicio - folga), min(len(texto), fim + folga)
    return {"antes": ("…" if a else "") + texto[a:inicio], "trecho": texto[inicio:fim], "depois": texto[fim:b] + ("…" if b < len(texto) else "")}


def checar(lista: list[dict]) -> list[dict]:
    """Achados do LanguageTool em cada trecho (posições relativas ao trecho)."""
    ignoradas = {p.lower() for p in dicionario()}
    achados, i = [], 0
    while i < len(lista):
        # monta um lote de trechos separados por linha em branco e guarda onde cada um começa
        inicios, partes, tamanho, j = [], [], 0, i
        while j < len(lista) and (not partes or tamanho + len(lista[j]["texto"]) + 2 <= LOTE):
            inicios.append(tamanho)
            partes.append(lista[j]["texto"])
            tamanho += len(lista[j]["texto"]) + 2
            j += 1
        for m in _consultar("\n\n".join(partes)) if any(p.strip() for p in partes) else []:
            k = max(n for n, ini in enumerate(inicios) if ini <= m["offset"])
            inicio, fim = m["offset"] - inicios[k], m["offset"] - inicios[k] + m["length"]
            texto = partes[k]
            if fim > len(texto) or m["rule"]["id"] in IGNORAR_REGRAS:
                continue
            original = texto[inicio:fim]
            if original.strip().lower() in ignoradas:
                continue
            achados.append({
                "id": f"a{len(achados) + 1}",
                "trecho": i + k,
                "rotulo": lista[i + k]["rotulo"],
                "inicio": inicio,
                "tamanho": fim - inicio,
                "original": original,
                "sugestoes": [r["value"] for r in m["replacements"][:5]],
                "mensagem": m["message"],
                "categoria": m["rule"]["category"]["name"],
                "regra": m["rule"]["id"],
                "contexto": _contexto(texto, inicio, fim),
            })
        i = j
    return achados


# ---------- Fluxo da revisão ----------


def _material(material_id: str) -> dict:
    m = materiais.obter(material_id)
    if m is None:
        raise LookupError("Material não encontrado")
    return m


def _formatado(m: dict) -> tuple[bytes, str]:
    achado = materiais.arquivo(m, "formatado")
    if achado is None:
        raise ValueError("O arquivo formatado não foi encontrado.")
    caminho = achado[0]
    return caminho.read_bytes(), caminho.suffix


def analisar(material_id: str, autor: dict) -> dict:
    m = _material(material_id)
    if m.get("etapa") != materiais.REVISAO:
        raise ValueError("Este material não está na revisão.")
    conteudo, ext = _formatado(m)
    lista = trechos(conteudo, ext)
    achados = checar(lista)
    revisao = {
        **(m.get("revisao") or {}),
        "analisado_em": agora().isoformat(timespec="minutes"),
        "analisado_por": autor["nome"],
        "formato": ext,
        "trechos": [t["texto"] for t in lista],
        "palavras": sum(len(t["texto"].split()) for t in lista),
        "achados": achados,
    }
    estado.alterar("materiais", material_id, {"revisao": revisao})
    return detalhe(material_id)


def detalhe(material_id: str) -> dict:
    m = _material(material_id)
    revisao = m.get("revisao") or {}
    publico = materiais._publico(m)
    publico["revisao"] = {k: v for k, v in revisao.items() if not k.startswith("arquivo") and k != "trechos"}
    publico["revisao"]["achados"] = revisao.get("achados")
    return publico


# ---------- Word com "Controlar alterações" ----------


class _Marcador:
    """Aplica uma troca num parágrafo como alteração controlada (w:del + w:ins), sem perder a formatação."""

    def __init__(self, autor: str):
        self.autor = autor
        self.data = datetime.now().strftime("%Y-%m-%dT%H:%M:%SZ")
        self.proximo_id = 9000

    def _marca(self, tag: str):
        el = OxmlElement(tag)
        el.set(qn("w:id"), str(self.proximo_id))
        el.set(qn("w:author"), self.autor)
        el.set(qn("w:date"), self.data)
        self.proximo_id += 1
        return el

    @staticmethod
    def _posicoes(p: Paragraph) -> list[tuple]:
        lista, pos = [], 0
        for r in p._p.iterchildren(qn("w:r")):
            texto = _texto_run(r)
            so_texto = all(x.tag in (qn("w:t"), qn("w:rPr")) for x in r)
            lista.append((r, pos, pos + len(texto), so_texto))
            pos += len(texto)
        return lista

    def _dividir(self, p: Paragraph, posicao: int):
        for r, ini, fim, simples in self._posicoes(p):
            if ini < posicao < fim:
                if not simples:
                    raise ValueError("trecho com tabulação ou quebra no meio")
                ts = list(r.iter(qn("w:t")))
                texto = "".join(t.text or "" for t in ts)
                for t in ts[1:]:
                    t.getparent().remove(t)
                novo = copy.deepcopy(r)
                for alvo, parte in ((ts[0], texto[: posicao - ini]), (next(novo.iter(qn("w:t"))), texto[posicao - ini :])):
                    alvo.text = parte
                    alvo.set(qn("xml:space"), "preserve")
                r.addnext(novo)
                return

    def trocar(self, p: Paragraph, inicio: int, fim: int, novo: str):
        self._dividir(p, inicio)
        self._dividir(p, fim)
        alvo = [r for r, ini, f, _ in self._posicoes(p) if ini >= inicio and f <= fim and f > ini]
        if not alvo:
            raise ValueError("trecho não encontrado")
        apagado = self._marca("w:del")
        alvo[0].addprevious(apagado)
        for r in alvo:
            apagado.append(r)
            for t in r.iter(qn("w:t")):
                t.tag = qn("w:delText")
        if novo:
            inserido = self._marca("w:ins")
            run = OxmlElement("w:r")
            formato = alvo[0].find(qn("w:rPr"))
            if formato is not None:
                run.append(copy.deepcopy(formato))
            t = OxmlElement("w:t")
            t.text = novo
            t.set(qn("xml:space"), "preserve")
            run.append(t)
            inserido.append(run)
            apagado.addnext(inserido)


def gerar_corrigido(material_id: str, escolhas: dict[str, str], autor: dict) -> dict:
    """Word com as correções aceitas como alterações controladas. `escolhas` = {achado_id: texto novo}."""
    m = _material(material_id)
    revisao = m.get("revisao") or {}
    if not revisao.get("achados"):
        raise ValueError("Analise o texto antes de gerar o arquivo corrigido.")
    por_id = {a["id"]: a for a in revisao["achados"]}
    desconhecidos = set(escolhas) - set(por_id)
    if desconhecidos:
        raise ValueError("Correção que não está na análise: refaça a análise.")

    conteudo, ext = _formatado(m)
    if ext == ".docx":
        doc = Document(BytesIO(conteudo))
        paragrafos = [p for _, p in _paragrafos_docx(doc)]
        if [_texto(p) for p in paragrafos] != revisao["trechos"]:
            raise ValueError("O arquivo formatado mudou depois da análise: analise de novo.")
    else:
        doc = Document()
        doc.add_heading(f"Texto do PDF revisado · {m['nome_arquivo']}", level=2)
        paragrafos = [doc.add_paragraph(t) for t in revisao["trechos"]]

    marcador = _Marcador(autor["nome"])
    aplicadas, puladas = 0, []
    # do fim para o começo de cada parágrafo, para as posições anteriores não mudarem
    for a in sorted((por_id[i] for i in escolhas), key=lambda a: (a["trecho"], -a["inicio"])):
        try:
            marcador.trocar(paragrafos[a["trecho"]], a["inicio"], a["inicio"] + a["tamanho"], escolhas[a["id"]])
            aplicadas += 1
        except ValueError:
            puladas.append(a["id"])
    # o que o revisor mudar depois no Word também fica marcado
    doc.settings.element.append(OxmlElement("w:trackRevisions"))

    saida = BytesIO()
    doc.save(saida)
    nome = f"{Path(m['nome_arquivo']).stem} - revisado.docx"
    estado.alterar("materiais", material_id, {"revisao": {
        **revisao,
        "arquivo_corrigido": arquivos.gravar("materiais", ".docx", saida.getvalue()),
        "nome_corrigido": nome,
        "aplicadas": aplicadas,
        "puladas": puladas,
        "corrigido_em": agora().isoformat(timespec="minutes"),
    }})
    return {"aplicadas": aplicadas, "puladas": puladas, "nome": nome}


def concluir(material_id: str, autor: dict, observacao: str = "", nome_arquivo: str = "", conteudo: bytes = b"") -> dict:
    """Libera para impressão. O arquivo final é o enviado agora, o corrigido ou, sem correções, o formatado."""
    with estado.transacao():
        m = _material(material_id)
        if m.get("etapa") != materiais.REVISAO:
            raise ValueError("Este material não está na revisão.")
        revisao = dict(m.get("revisao") or {})
        if conteudo:
            ext = arquivos.validar_editavel(nome_arquivo, conteudo)
            revisao["arquivo_final"], revisao["nome_final"] = arquivos.gravar("materiais", ext, conteudo), Path(nome_arquivo).name
        elif revisao.get("arquivo_corrigido"):
            revisao["arquivo_final"], revisao["nome_final"] = revisao["arquivo_corrigido"], revisao["nome_corrigido"]
        else:
            revisao["arquivo_final"], revisao["nome_final"] = m["diagramacao"]["arquivo"], m["diagramacao"]["nome_arquivo"]
        revisao.update(concluida_por=autor["nome"], concluida_em=agora().isoformat(timespec="minutes"), observacao=observacao.strip())
        estado.alterar("materiais", material_id, {"etapa": materiais.PRONTO, "revisao": revisao})
    return detalhe(material_id)


def devolver(material_id: str, motivo: str, autor: dict) -> dict:
    """Volta para a diagramação (ex.: problema de formatação que a revisão não resolve)."""
    motivo = motivo.strip()
    if len(motivo) < 3:
        raise ValueError("Diga ao diagramador o que precisa ser refeito.")
    with estado.transacao():
        m = _material(material_id)
        if m.get("etapa") != materiais.REVISAO:
            raise ValueError("Este material não está na revisão.")
        estado.alterar("materiais", material_id, {
            "etapa": materiais.DIAGRAMACAO,
            "devolucao": {"motivo": motivo, "por": autor["nome"], "em": agora().isoformat(timespec="minutes")},
        })
    return detalhe(material_id)
