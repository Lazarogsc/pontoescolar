"""Termo de matrícula em .docx, para o responsável e a secretaria assinarem.

As cláusulas são um modelo genérico: a escola deve revisá-las com o jurídico
antes de usar em produção.
"""

from datetime import date, datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.util import br, MESES_LONGOS

MARINHO = RGBColor(0x0F, 0x2A, 0x6B)
SUAVE = RGBColor(0x55, 0x62, 0x7F)


def _moeda(v: float) -> str:
    inteiro, centavos = f"{v:,.2f}".split(".")
    return f"R$ {inteiro.replace(',', '.')},{centavos}"


def _data(iso: str | None) -> str:
    return br(date.fromisoformat(iso)) if iso else "—"


def _fundo(celula, cor: str) -> None:
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), cor)
    celula._tc.get_or_add_tcPr().append(sombra)


def _sem_bordas(tabela) -> None:
    bordas = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{lado}")
        el.set(qn("w:val"), "nil")
        bordas.append(el)
    tabela._tbl.tblPr.append(bordas)


def _titulo_secao(doc: Document, texto: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(texto.upper())
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = MARINHO


def _campos(doc: Document, pares: list[tuple[str, str]], colunas: int = 2) -> None:
    """Tabela de rótulo/valor, `colunas` pares por linha."""
    linhas = (len(pares) + colunas - 1) // colunas
    t = doc.add_table(rows=linhas, cols=colunas * 2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (rotulo, valor) in enumerate(pares):
        lin, col = divmod(i, colunas)
        c_rot, c_val = t.cell(lin, col * 2), t.cell(lin, col * 2 + 1)
        _fundo(c_rot, "EEF2FA")
        c_rot.text, c_val.text = rotulo, valor or "—"
        c_rot.paragraphs[0].runs[0].bold = True
        for c in (c_rot, c_val):
            c.paragraphs[0].runs[0].font.size = Pt(9)
    # Último par sozinho ocupa a linha inteira
    if len(pares) % colunas:
        ultima = t.rows[-1]
        ultima.cells[1].merge(ultima.cells[-1])


CLAUSULAS = [
    "O(A) CONTRATANTE matricula o(a) aluno(a) acima identificado(a) na série/turma indicada, para o ano letivo de {ano}, "
    "comprometendo-se a cumprir o Regimento Escolar e as normas internas da CONTRATADA, das quais declara ter ciência.",
    "Pelos serviços educacionais, o(a) CONTRATANTE pagará a anuidade de {total}, dividida em {parcelas} parcela(s) mensais "
    "de {valor}, com vencimento todo dia {dia}, sendo a primeira em {primeiro}.",
    "O pagamento após o vencimento sujeita o(a) CONTRATANTE à multa de 2% (dois por cento) e a juros de mora de 1% (um por cento) "
    "ao mês, calculados pró-rata dia. O atraso caracteriza inadimplência e poderá gerar cobrança extrajudicial.",
    "Material didático, uniforme, atividades extracurriculares, passeios e eventos não estão incluídos no valor da anuidade.",
    "O cancelamento da matrícula deverá ser solicitado por escrito à Secretaria, com quitação das parcelas vencidas até a data do pedido.",
    "O(A) CONTRATANTE declara serem verdadeiras as informações prestadas e autoriza o tratamento dos dados pessoais do(a) aluno(a) "
    "e do responsável para fins escolares, nos termos da Lei nº 13.709/2018 (LGPD).",
]


def gerar(m: dict) -> bytes:
    """`m` é o detalhe da matrícula (escola.detalhe_matricula)."""
    escola, aluno, resp, turma, contrato = m["escola"], m["aluno"], m["responsavel_detalhe"], m["turma_detalhe"], m["contrato"]
    doc = Document()

    secao = doc.sections[0]
    secao.page_height, secao.page_width = Cm(29.7), Cm(21)
    secao.top_margin = secao.bottom_margin = Cm(1.5)
    secao.left_margin = secao.right_margin = Cm(2)

    estilo = doc.styles["Normal"]
    estilo.font.name = "Calibri"
    estilo.font.size = Pt(10)
    estilo.paragraph_format.space_after = Pt(2)

    # Cabeçalho da escola
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(escola["nome"].upper())
    r.bold, r.font.size, r.font.color.rgb = True, Pt(14), MARINHO
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"CNPJ {escola['cnpj']} · {escola['endereco']}")
    r.font.size, r.font.color.rgb = Pt(8.5), SUAVE

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run(f"TERMO DE MATRÍCULA E CONTRATO DE PRESTAÇÃO DE SERVIÇOS EDUCACIONAIS — {m['ano_letivo']}")
    r.bold, r.font.size = True, Pt(11.5)

    # Número da matrícula em destaque
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _sem_bordas(t)
    for celula, (rotulo, valor) in zip(t.rows[0].cells, [
        ("Nº DA MATRÍCULA", m["matricula"]),
        ("DATA DA MATRÍCULA", _data(m["data_matricula"])),
        ("TIPO", m["tipo"]),
    ]):
        _fundo(celula, "EAF0FF")
        celula.text = ""
        p1 = celula.paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p1.add_run(rotulo)
        r1.font.size, r1.font.color.rgb, r1.bold = Pt(7.5), SUAVE, True
        p2 = celula.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(valor)
        r2.bold, r2.font.size, r2.font.color.rgb = True, Pt(13), MARINHO

    _titulo_secao(doc, "1. Dados do(a) aluno(a)")
    _campos(doc, [
        ("Nome", aluno["nome"]),
        ("Nascimento", _data(aluno["data_nascimento"])),
        ("CPF", aluno.get("cpf") or "—"),
        ("Sexo", aluno.get("sexo") or "—"),
        ("Escola anterior", aluno.get("escola_anterior") or "—"),
    ])

    _titulo_secao(doc, "2. Responsável financeiro (CONTRATANTE)")
    _campos(doc, [
        ("Nome", resp["nome"]),
        ("CPF", resp["cpf"]),
        ("Parentesco", resp["parentesco"]),
        ("Telefone", resp["telefone"]),
        ("E-mail", resp["email"]),
        ("Endereço", resp["endereco"]),
    ])

    _titulo_secao(doc, "3. Curso e turma")
    _campos(doc, [
        ("Segmento", turma["segmento"]),
        ("Turma", turma["nome"]),
        ("Turno", turma["turno"]),
        ("Sala", turma["sala"]),
    ])

    _titulo_secao(doc, "4. Condições financeiras")
    desconto = f"{contrato['desconto']}%" if contrato["desconto"] else "Sem desconto"
    _campos(doc, [
        ("Mensalidade", _moeda(contrato["valor_bruto"])),
        ("Desconto", desconto),
        ("Valor da parcela", _moeda(contrato["valor_mensal"])),
        ("Parcelas", f"{contrato['parcelas']}x"),
        ("Vencimento", f"Todo dia {contrato['dia_vencimento']}"),
        ("Total do contrato", _moeda(contrato["total"])),
    ])

    if aluno.get("observacoes"):
        _titulo_secao(doc, "Observações")
        doc.add_paragraph(aluno["observacoes"])

    _titulo_secao(doc, "5. Cláusulas")
    valores = {
        "ano": m["ano_letivo"], "total": _moeda(contrato["total"]), "parcelas": contrato["parcelas"],
        "valor": _moeda(contrato["valor_mensal"]), "dia": contrato["dia_vencimento"], "primeiro": _data(contrato["primeiro_vencimento"]),
    }
    for i, clausula in enumerate(CLAUSULAS, start=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"Cláusula {i}ª. ")
        r.bold = True
        p.add_run(clausula.format(**valores)).font.size = Pt(9.5)

    # Local e data
    criada = date.fromisoformat(m["data_matricula"])
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(14)
    p.add_run(f"{escola['cidade']}, {criada.day} de {MESES_LONGOS[criada.month - 1]} de {criada.year}.")

    # Assinaturas lado a lado
    t = doc.add_table(rows=1, cols=2)
    _sem_bordas(t)
    assinaturas = [
        ("Assinatura do responsável financeiro", resp["nome"], f"CPF {resp['cpf']}"),
        ("Assinatura de quem realizou a matrícula", m["atendente"], f"Secretaria · {escola['nome']}"),
    ]
    for celula, (papel, nome, detalhe) in zip(t.rows[0].cells, assinaturas):
        celula.text = ""
        p = celula.paragraphs[0]
        p.paragraph_format.space_before = Pt(34)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run("_" * 38)
        for texto, negrito, tamanho in ((nome, True, 10), (detalhe, False, 8.5), (papel, False, 8.5)):
            q = celula.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            q.paragraph_format.space_after = Pt(0)
            run = q.add_run(texto)
            run.bold, run.font.size = negrito, Pt(tamanho)
            if not negrito:
                run.font.color.rgb = SUAVE

    # O Word exige um parágrafo depois da última tabela; mínimo, para não abrir uma página em branco.
    fim = doc.add_paragraph()
    fim.paragraph_format.space_before = fim.paragraph_format.space_after = Pt(0)
    fim.add_run().font.size = Pt(1)

    # Rodapé
    rodape = secao.footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = rodape.add_run(f"Matrícula nº {m['matricula']} · Documento gerado pelo ponto.escolar em {datetime.now():%d/%m/%Y %H:%M}")
    r.font.size, r.font.color.rgb = Pt(7.5), SUAVE

    saida = BytesIO()
    doc.save(saida)
    return saida.getvalue()
