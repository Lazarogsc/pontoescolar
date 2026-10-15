"""Diagramação e revisão. O LanguageTool é trocado por um corretor falso com erros conhecidos:
o que se testa aqui é o fluxo, a extração do texto e o Word com alterações controladas."""

import re
from datetime import timedelta
from io import BytesIO

import pytest
from docx import Document
from docx.oxml.ns import qn
from reportlab.pdfgen.canvas import Canvas

from app import revisao
from app.util import hoje
from tests.apoio import entrar_como

ERROS = {"exercicios": ["exercícios"], "Os alunos deve": ["Os alunos devem"], "area": ["área", "áreas"], "Agente": ["A gente"]}


def _falso_languagetool(texto: str) -> list[dict]:
    achados = []
    for errado, sugestoes in ERROS.items():
        for m in re.finditer(rf"\b{re.escape(errado)}\b", texto):
            achados.append({"offset": m.start(), "length": len(errado), "message": f"Possível erro: {errado}",
                            "replacements": [{"value": s} for s in sugestoes],
                            "rule": {"id": "FALSO", "category": {"name": "Gramática"}}})
    return achados


@pytest.fixture(autouse=True)
def corretor(monkeypatch):
    monkeypatch.setattr(revisao, "_consultar", _falso_languagetool)


def _docx() -> bytes:
    doc = Document()
    doc.add_heading("Prova de Matemática", level=1)
    p = doc.add_paragraph()
    p.add_run("Os alunos ").bold = True  # o erro atravessa dois trechos com formatação diferente
    p.add_run("deve resolver os exercicios com calma.")
    doc.add_paragraph("1) Calcule a area do triângulo.")
    tabela = doc.add_table(rows=1, cols=2)
    tabela.cell(0, 0).text = "Agente confere no fim."
    tabela.cell(0, 1).text = "Sem erros aqui."
    saida = BytesIO()
    doc.save(saida)
    return saida.getvalue()


def _pdf() -> bytes:
    saida = BytesIO()
    c = Canvas(saida)
    for i, linha in enumerate(["Resolva os exercicios da lista com", "atenção e calcule a area de cada fi-", "gura.", "", "2) Agente corrige depois."]):
        c.drawString(72, 750 - i * 16, linha)
    c.save()
    return saida.getvalue()


def _texto_aceito(conteudo: bytes) -> str:
    """Texto do documento como fica depois de 'Aceitar todas as alterações' no Word."""
    corpo = Document(BytesIO(conteudo)).element.body
    return "".join(t.text or "" for t in corpo.iter(qn("w:t")))


def _material_aprovado(prof, coord, tipo="Prova"):
    data = (hoje() + timedelta(days=10)).isoformat()
    m = prof.post("/api/professores/materiais", data={"data": data, "turma": "6a", "tipo": tipo},
                  files={"arquivo": ("prova.pdf", b"%PDF-1.4 x")}).json()
    coord.post(f"/api/cpd/materiais/{m['id']}/avaliacao", data={"status": "Aprovado"})
    return m


def test_do_professor_a_impressao_com_word_corrigido():
    prof, coord = entrar_como("prof.rev", "Docentes"), entrar_como("coord.rev", "Coordenação do CPD")
    diag, rev = entrar_como("diag.rev", "Diagramação"), entrar_como("rev.rev", "Revisão")
    m = _material_aprovado(prof, coord)
    assert m["tipo"] == "Prova"

    pendentes = diag.get("/api/cpd/diagramacao").json()["itens"]
    assert any(i["id"] == m["id"] and i["etapa"] == "Diagramação" for i in pendentes)
    assert diag.get(f"/api/cpd/materiais/{m['id']}/arquivo").status_code == 200  # baixa o original para formatar
    assert "Word" in diag.post(f"/api/cpd/diagramacao/{m['id']}", files={"arquivo": ("x.png", b"\x89PNG")}).json()["detail"]
    r = diag.post(f"/api/cpd/diagramacao/{m['id']}", files={"arquivo": ("prova-formatada.docx", _docx())}, data={"observacao": "Fonte 12"})
    assert r.status_code == 200 and r.json()["etapa"] == "Revisão"
    assert any(i["id"] == m["id"] for i in diag.get("/api/cpd/diagramacao", params={"aba": "formatados"}).json()["itens"])
    assert {i["id"]: i for i in prof.get("/api/professores/materiais").json()["itens"]}[m["id"]]["etapa"] == "Revisão"

    analise = rev.post(f"/api/cpd/revisao/{m['id']}/analise").json()
    achados = analise["revisao"]["achados"]
    assert sorted(a["original"] for a in achados) == ["Agente", "Os alunos deve", "area", "exercicios"]
    assert {a["rotulo"] for a in achados if a["original"] == "Agente"} == {"Tabela 1, linha 1"}

    # aceita três, recusa "area"; em "Os alunos deve" escolhe a sugestão
    escolhas = {a["id"]: a["sugestoes"][0] for a in achados if a["original"] != "area"}
    gerado = rev.post(f"/api/cpd/revisao/{m['id']}/corrigido", json={"escolhas": escolhas}).json()
    assert gerado["aplicadas"] == 3 and gerado["puladas"] == []

    corrigido = rev.get(f"/api/cpd/materiais/{m['id']}/arquivo", params={"versao": "corrigido"}).content
    corpo = Document(BytesIO(corrigido)).element.body
    assert len(list(corpo.iter(qn("w:ins")))) == 3 and len(list(corpo.iter(qn("w:del")))) == 3
    texto = _texto_aceito(corrigido)
    assert "Os alunos devem resolver os exercícios com calma." in texto
    assert "Calcule a area do triângulo." in texto and "A gente confere no fim." in texto
    # o que entrou no lugar de "Os alunos deve" herdou o negrito do começo do trecho
    novo = next(corpo.iter(qn("w:ins"))).find(qn("w:r"))
    assert novo.find(qn("w:rPr")) is not None

    final = rev.post(f"/api/cpd/revisao/{m['id']}/conclusao", data={"observacao": "Ok"}).json()
    assert final["etapa"] == "Pronto para impressão" and final["revisao"]["nome_final"].endswith("revisado.docx")
    assert rev.get(f"/api/cpd/materiais/{m['id']}/arquivo", params={"versao": "final"}).content == corrigido
    assert {i["id"]: i for i in prof.get("/api/professores/materiais").json()["itens"]}[m["id"]]["etapa"] == "Pronto para impressão"


def test_pdf_vira_word_com_as_correcoes():
    prof, coord = entrar_como("prof.pdf", "Docentes"), entrar_como("coord.pdf", "Coordenação do CPD")
    diag, rev = entrar_como("diag.pdf", "Diagramação"), entrar_como("rev.pdf", "Revisão")
    m = _material_aprovado(prof, coord, tipo="Lista de exercícios")
    diag.post(f"/api/cpd/diagramacao/{m['id']}", files={"arquivo": ("lista.pdf", _pdf())})
    achados = rev.post(f"/api/cpd/revisao/{m['id']}/analise").json()["revisao"]["achados"]
    # as linhas quebradas viraram um parágrafo só e "fi-/gura" foi juntado
    assert sorted(a["original"] for a in achados) == ["Agente", "area", "exercicios"]
    rev.post(f"/api/cpd/revisao/{m['id']}/corrigido", json={"escolhas": {a["id"]: a["sugestoes"][0] for a in achados}})
    texto = _texto_aceito(rev.get(f"/api/cpd/materiais/{m['id']}/arquivo", params={"versao": "corrigido"}).content)
    assert "Resolva os exercícios da lista com atenção e calcule a área de cada figura." in texto
    assert "2) A gente corrige depois." in texto


def test_devolucao_dicionario_e_permissoes():
    prof, coord = entrar_como("prof.dev", "Docentes"), entrar_como("coord.dev", "Coordenação do CPD")
    diag, rev = entrar_como("diag.dev", "Diagramação"), entrar_como("rev.dev", "Revisão")
    m = _material_aprovado(prof, coord)
    assert "não está na revisão" in rev.post(f"/api/cpd/revisao/{m['id']}/analise").json()["detail"]
    diag.post(f"/api/cpd/diagramacao/{m['id']}", files={"arquivo": ("f.docx", _docx())})

    assert "diagramador" in rev.post(f"/api/cpd/revisao/{m['id']}/devolucao", data={"motivo": ""}).json()["detail"]
    volta = rev.post(f"/api/cpd/revisao/{m['id']}/devolucao", data={"motivo": "Tabela cortada na página 2"}).json()
    assert volta["etapa"] == "Diagramação" and volta["devolucao"]["motivo"] == "Tabela cortada na página 2"
    diag.post(f"/api/cpd/diagramacao/{m['id']}", files={"arquivo": ("f2.docx", _docx())})

    # "Agente" é nome próprio nesta escola: vai para o dicionário e some da análise
    rev.post("/api/cpd/revisao/dicionario", json={"palavra": "Agente"})
    achados = rev.post(f"/api/cpd/revisao/{m['id']}/analise").json()["revisao"]["achados"]
    assert "Agente" not in [a["original"] for a in achados] and "Agente" in rev.get("/api/cpd/revisao/dicionario").json()

    assert diag.get("/api/cpd/revisao").status_code == 403
    assert rev.get("/api/cpd/diagramacao").status_code == 403
    assert prof.get("/api/cpd/diagramacao").status_code == 403
    assert rev.get("/api/cpd/materiais").status_code == 403  # recebimento é da equipe do CPD
    assert coord.get("/api/cpd/revisao").status_code == 200 and coord.get("/api/cpd/diagramacao").status_code == 200


def test_corretor_fora_do_ar(monkeypatch):
    monkeypatch.undo()
    monkeypatch.setattr(revisao, "URL", "http://127.0.0.1:9")
    prof, coord = entrar_como("prof.fora", "Docentes"), entrar_como("coord.fora", "Coordenação do CPD")
    diag, rev = entrar_como("diag.fora", "Diagramação"), entrar_como("rev.fora", "Revisão")
    m = _material_aprovado(prof, coord)
    diag.post(f"/api/cpd/diagramacao/{m['id']}", files={"arquivo": ("f.docx", _docx())})
    r = rev.post(f"/api/cpd/revisao/{m['id']}/analise")
    assert r.status_code == 503 and "LanguageTool" in r.json()["detail"]
